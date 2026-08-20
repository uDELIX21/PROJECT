"""Bulk import engine (REQ-IMP-*, BR-I, design §10).

Pipeline: Upload → Parse → Validate → Preview → Confirm → Import → Report.
Commit is transaction-safe: everything lands or nothing does (BR-I01).
Parent links are only SUGGESTED by phone match — a human confirms each one
(Agent Rule 8)."""
import csv
import io
import uuid
from datetime import date, datetime

from fastapi import Request
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.audit import audit
from app.core.errors import ConflictError, NotFoundError
from app.core.ids import uuid7
from app.core.phones import normalize_ghana_phone
from app.models.base import utcnow
from app.models.core import (ClassStream, Enrollment, Grade, ParentGuardian,
                             ParentStudentRelationship, Student, Subject, Teacher,
                             TeacherAssignment)
from app.models.imports import ImportErrorRow, ImportJob
from app.services import students as students_svc
from app.services import guardians as guardians_svc
from app.services import enrollments as enroll_svc
from app.services import fees as fees_svc
from app.services import ledger

# --------------------------------------------------------------------------- templates
# Single source for downloads AND validation (drift-proof, design §10).

TEMPLATES: dict[str, dict] = {
    "STUDENTS": {
        "columns": ["admission_code", "surname", "other_names", "gender", "date_of_birth",
                    "class", "guardian_name", "guardian_relationship", "guardian_phone",
                    "guardian_email", "guardian_address", "ghana_digital_address",
                    "opening_balance_ghs"],
        "required": ["surname", "other_names", "gender", "date_of_birth", "class"],
        "example": ["", "Mensah", "Kofi", "M", "2018-04-12", "Basic 4A",
                    "Mensah Abena", "MOTHER", "0241234567", "parent@example.com",
                    "12 Palm Street, Accra", "GA-123-4567", "150.00"],
    },
    "TEACHERS": {
        "columns": ["staff_code", "surname", "other_names", "gender", "phone", "email",
                    "job_title", "qualification", "hired_on", "primary_class",
                    "assigned_subjects"],
        "required": ["surname", "other_names"],
        "example": ["T-101", "Owusu", "Daniel", "M", "0209876543", "d.owusu@example.com",
                    "Teacher", "B.Ed Mathematics", "2022-09-01", "Basic 4A", "MAT;SCI"],
    },
    "PARENTS": {
        "columns": ["name", "relationship_to", "student_admission_code", "student_name",
                    "phone", "phone2", "email", "occupation", "residential_address",
                    "ghana_digital_address", "latitude", "longitude",
                    "preferred_channel", "contact_window"],
        "required": ["name", "phone"],
        "example": ["Boateng Ama", "MOTHER", "HSA-0001", "", "0244567890", "",
                    "ama@example.com", "Trader", "8 School Road, Kumasi", "AK-334-5566",
                    "", "", "SMS", "17:00-20:00"],
    },
}

GENDER_ALIASES = {"m": "M", "male": "M", "f": "F", "female": "F"}
REL_ALIASES = {"mother": "MOTHER", "father": "FATHER", "guardian": "GUARDIAN",
               "grandparent": "GRANDPARENT", "grandmother": "GRANDPARENT",
               "grandfather": "GRANDPARENT", "sibling": "SIBLING", "other": "OTHER"}


def template_csv(kind: str) -> bytes:
    t = TEMPLATES[kind]
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(t["columns"])
    w.writerow(t["example"])
    return out.getvalue().encode("utf-8-sig")


def template_xlsx(kind: str) -> bytes:
    from openpyxl import Workbook
    t = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = kind.title()
    ws.append(t["columns"])
    ws.append(t["example"])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Field notes"])
    ws2.append(["gender: M/F (or Male/Female)"])
    ws2.append(["date_of_birth / hired_on: YYYY-MM-DD or DD/MM/YYYY"])
    ws2.append(["class: stream name for the ACTIVE academic year, e.g. 'Basic 4A'"])
    ws2.append(["opening_balance_ghs: decimal GHS owed at import time (optional)"])
    ws2.append(["assigned_subjects: subject codes separated by ';'"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- parsing

def parse_file(kind: str, data: bytes, filename: str) -> list[dict]:
    """CSV (utf-8-sig/latin-1) or XLSX → list of row dicts keyed by header."""
    if kind not in TEMPLATES:
        raise ConflictError(f"Unknown import kind {kind}.", code="KIND_UNKNOWN")
    rows: list[dict] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ConflictError("Could not read the workbook (corrupt XLSX?).",
                                code="FILE_UNREADABLE")
        ws = wb[wb.sheetnames[0]]
        header = None
        for raw in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip().lower() if c is not None else "" for c in raw]
                continue
            if all(c is None for c in raw):
                continue
            rows.append({header[i]: (raw[i] if i < len(raw) else None)
                         for i in range(len(header)) if header[i]})
    else:
        text = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ConflictError("File encoding not recognized.", code="FILE_UNREADABLE")
        reader = csv.DictReader(io.StringIO(text))
        if reader.fieldnames is None:
            raise ConflictError("CSV has no header row.", code="FILE_UNREADABLE")
        for raw in reader:
            rows.append({(k or "").strip().lower(): (v.strip() if isinstance(v, str) else v)
                         for k, v in raw.items()})
    return rows


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_amount_ghs(value) -> int | None:
    """GHS string → pesewas. Returns None when blank; raises ValueError when bad."""
    if value in (None, ""):
        return None
    s = str(value).replace(",", "").replace("GH¢", "").replace("GHS", "").strip()
    try:
        pesewas = round(float(s) * 100)
    except ValueError:
        raise ValueError(f"'{value}' is not a valid amount")
    if pesewas < 0:
        raise ValueError("amount must not be negative")
    return pesewas


# --------------------------------------------------------------------------- validation

def _streams_by_name(db: Session, school_id: uuid.UUID, year_id: uuid.UUID) -> dict[str, ClassStream]:
    rows = db.scalars(select(ClassStream).where(
        ClassStream.school_id == school_id, ClassStream.academic_year_id == year_id)).all()
    return {r.name.strip().lower(): r for r in rows}


def _nearest_stream(name: str, known: dict[str, ClassStream]) -> str | None:
    """Crude 'did you mean' suggestion: first token (grade) match."""
    first = name.split()[0] if name.split() else ""
    for key in known:
        if key.split()[0] == first:
            return known[key].name
    return None


def _guardian_matches(db: Session, school_id: uuid.UUID, phone_e164: str,
                      surname_hint: str | None) -> list[dict]:
    """Phone-based CANDIDATES only — never auto-linked (Agent Rule 8)."""
    q = select(ParentGuardian).where(ParentGuardian.school_id == school_id,
                                     ParentGuardian.phone == phone_e164)
    out = []
    for g in db.scalars(q).all():
        confidence = "high"
        if surname_hint and surname_hint.strip().lower() in g.name.lower():
            confidence = "very high"
        out.append({"guardian_id": str(g.id), "guardian_name": g.name,
                    "phone": g.phone, "confidence": confidence})
    return out


def validate_rows(db: Session, *, school_id: uuid.UUID, kind: str, rows: list[dict],
                  year_id: uuid.UUID) -> list[dict]:
    """Per-row findings. Each finding: row_no, severity, field, message, guidance,
    match_suggestion."""
    tpl = TEMPLATES[kind]
    streams = _streams_by_name(db, school_id, year_id)
    findings: list[dict] = []
    seen_codes: dict[str, int] = {}
    seen_names: dict[str, int] = {}

    for idx, row in enumerate(rows, start=2):  # row 1 = header
        issues: list[dict] = []
        norm = {(k or "").strip().lower(): (v if v is not None else "") for k, v in row.items()}

        # required fields
        for col in tpl["required"]:
            if not str(norm.get(col, "")).strip():
                issues.append({"severity": "ERROR", "field": col,
                               "message": f"'{col}' is required.",
                               "guidance": "Fill in this column and re-upload."})

        if kind == "STUDENTS":
            code = str(norm.get("admission_code", "")).strip()
            gender = str(norm.get("gender", "")).strip().lower()
            dob = _parse_date(norm.get("date_of_birth"))
            cls = str(norm.get("class", "")).strip()
            gphone = str(norm.get("guardian_phone", "")).strip()
            opening = norm.get("opening_balance_ghs")

            if code:
                key = code.lower()
                if key in seen_codes:
                    issues.append({"severity": "DUPLICATE", "field": "admission_code",
                                   "message": f"Admission code '{code}' appears more than once in this file (rows {seen_codes[key]} and {idx}).",
                                   "guidance": "Keep one row per student."})
                else:
                    seen_codes[key] = idx
                    exists = db.scalar(select(Student).where(
                        Student.school_id == school_id, Student.admission_code == code))
                    if exists is not None:
                        issues.append({"severity": "DUPLICATE", "field": "admission_code",
                                       "message": f"'{code}' already exists ({exists.full_name}).",
                                       "guidance": "Remove this row or clear the admission_code to create a new student."})
            name_key = f"{norm.get('surname', '')}|{norm.get('other_names', '')}|{dob}".lower()
            if str(norm.get("surname", "")).strip():
                if name_key in seen_names:
                    issues.append({"severity": "DUPLICATE", "field": "name",
                                   "message": f"Same name+DoB as row {seen_names[name_key]}.",
                                   "guidance": "Likely duplicate — keep one row."})
                else:
                    seen_names[name_key] = idx
                dup = db.scalar(select(Student).where(
                    Student.school_id == school_id,
                    Student.surname == str(norm.get("surname", "")).strip(),
                    Student.other_names == str(norm.get("other_names", "")).strip(),
                    Student.date_of_birth == dob)) if dob else None
                if dup is not None:
                    issues.append({"severity": "DUPLICATE", "field": "name",
                                   "message": f"Matches existing student {dup.admission_code}.",
                                   "guidance": "Check the register before importing again."})
            if gender and gender not in GENDER_ALIASES:
                issues.append({"severity": "ERROR", "field": "gender",
                               "message": f"Unknown gender '{norm.get('gender')}'.",
                               "guidance": "Use M/F (or Male/Female)."})
            if str(norm.get("date_of_birth", "")).strip() and dob is None:
                issues.append({"severity": "ERROR", "field": "date_of_birth",
                               "message": f"Unrecognized date '{norm.get('date_of_birth')}'.",
                               "guidance": "Use YYYY-MM-DD or DD/MM/YYYY."})
            if cls:
                stream = streams.get(cls.lower())
                if stream is None:
                    suggest = _nearest_stream(cls.lower(), streams)
                    issues.append({"severity": "ERROR", "field": "class",
                                   "message": f"No class '{cls}' in the active year.",
                                   "guidance": f"Did you mean '{suggest}'?" if suggest
                                   else "Use exact stream names, e.g. 'Basic 4A'."})
            if gphone:
                if normalize_ghana_phone(gphone) is None:
                    issues.append({"severity": "ERROR", "field": "guardian_phone",
                                   "message": f"'{gphone}' is not a valid Ghana number.",
                                   "guidance": "Use e.g. 0241234567 or +233241234567."})
                else:
                    matches = _guardian_matches(db, school_id, normalize_ghana_phone(gphone),
                                                str(norm.get("surname", "")))
                    if matches:
                        issues.append({"severity": "MATCH_CANDIDATE",
                                       "field": "guardian_phone",
                                       "message": f"Possible existing guardian: {matches[0]['guardian_name']} ({matches[0]['confidence']} confidence).",
                                       "guidance": "Confirm the link, choose to create a new guardian, or skip the link.",
                                       "match": {"candidates": matches}})
            if opening not in (None, ""):
                try:
                    _parse_amount_ghs(opening)
                except ValueError as e:
                    issues.append({"severity": "ERROR", "field": "opening_balance_ghs",
                                   "message": str(e),
                                   "guidance": "Use a decimal GHS amount, e.g. 150.00"})

        elif kind == "TEACHERS":
            phone = str(norm.get("phone", "")).strip()
            if phone and normalize_ghana_phone(phone) is None:
                issues.append({"severity": "ERROR", "field": "phone",
                               "message": f"'{phone}' is not a valid Ghana number.",
                               "guidance": "Use e.g. 0209876543."})
            if str(norm.get("hired_on", "")).strip() and _parse_date(norm.get("hired_on")) is None:
                issues.append({"severity": "ERROR", "field": "hired_on",
                               "message": "Unrecognized date.",
                               "guidance": "Use YYYY-MM-DD."})
            if email := str(norm.get("email", "")).strip():
                if "@" not in email:
                    issues.append({"severity": "WARNING", "field": "email",
                                   "message": f"'{email}' does not look like an email.",
                                   "guidance": "Check the address."})

        elif kind == "PARENTS":
            phone = str(norm.get("phone", "")).strip()
            if phone and normalize_ghana_phone(phone) is None:
                issues.append({"severity": "ERROR", "field": "phone",
                               "message": f"'{phone}' is not a valid Ghana number.",
                               "guidance": "Use e.g. 0244567890."})
            if phone and normalize_ghana_phone(phone):
                matches = _guardian_matches(db, school_id, normalize_ghana_phone(phone),
                                            str(norm.get("name", "")).split()[-1] if norm.get("name") else None)
                if matches:
                    issues.append({"severity": "MATCH_CANDIDATE", "field": "phone",
                                   "message": f"Possible existing guardian: {matches[0]['guardian_name']} ({matches[0]['confidence']} confidence).",
                                   "guidance": "Confirm the link or create a new guardian.",
                                   "match": {"candidates": matches}})
            code = str(norm.get("student_admission_code", "")).strip()
            if code:
                st = db.scalar(select(Student).where(Student.school_id == school_id,
                                                     Student.admission_code == code))
                if st is None:
                    issues.append({"severity": "WARNING", "field": "student_admission_code",
                                   "message": f"No student with code '{code}'.",
                                   "guidance": "The link will be skipped unless the student exists."})

        if not issues:
            findings.append({"row_no": idx, "severity": "OK", "message": "",
                             "raw": norm})
        else:
            worst = "OK"
            order = {"OK": 0, "WARNING": 1, "MATCH_CANDIDATE": 2, "DUPLICATE": 3, "ERROR": 4}
            for i in issues:
                if order[i["severity"]] > order[worst]:
                    worst = i["severity"]
            findings.append({"row_no": idx, "severity": worst, "issues": issues,
                             "raw": norm})
    return findings


# --------------------------------------------------------------------------- commit

def _import_student_row(db: Session, *, school_id: uuid.UUID, norm: dict,
                        year_id: uuid.UUID, actor_id: uuid.UUID,
                        match_resolution: dict | None, request: Request | None) -> str:
    streams = _streams_by_name(db, school_id, year_id)
    stream = streams[str(norm.get("class", "")).strip().lower()]
    student = students_svc.create_student(
        db, school_id=school_id, surname=str(norm["surname"]).strip(),
        other_names=str(norm["other_names"]).strip(),
        gender=GENDER_ALIASES[str(norm["gender"]).strip().lower()],
        date_of_birth=_parse_date(norm["date_of_birth"]), admit=True,
        actor_id=actor_id, request=request)
    enroll_svc.create_enrollment(db, school_id=school_id, student_id=student.id,
                                 class_stream_id=stream.id, actor_id=actor_id,
                                 request=request)
    # guardian handling
    gname = str(norm.get("guardian_name", "")).strip()
    gphone = str(norm.get("guardian_phone", "")).strip()
    guardian = None
    if match_resolution and match_resolution.get("action") == "LINK":
        guardian = db.get(ParentGuardian, uuid.UUID(match_resolution["guardian_id"]))
    elif match_resolution and match_resolution.get("action") == "SKIP":
        guardian = None
    elif gname and gphone:
        guardian = guardians_svc.create_guardian(
            db, school_id=school_id, name=gname, phone=gphone,
            email=str(norm.get("guardian_email", "")).strip() or None,
            residential_address=str(norm.get("guardian_address", "")).strip() or None,
            ghana_digital_address=str(norm.get("ghana_digital_address", "")).strip() or None,
            actor_id=actor_id, request=request)
    if guardian is not None:
        rel_type = REL_ALIASES.get(str(norm.get("guardian_relationship", "")).strip().lower(),
                                   "GUARDIAN")
        guardians_svc.link_student(db, school_id=school_id, parent_id=guardian.id,
                                   student_id=student.id, relationship_type=rel_type,
                                   actor_id=actor_id,
                                   source="IMPORT_SUGGESTION" if match_resolution else "IMPORT",
                                   request=request)
    # opening balance
    opening = norm.get("opening_balance_ghs")
    if opening not in (None, ""):
        pesewas = _parse_amount_ghs(opening)
        if pesewas:
            enrollment = db.scalar(select(Enrollment).where(
                Enrollment.student_id == student.id,
                Enrollment.academic_year_id == year_id,
                Enrollment.status == "ACTIVE"))
            fees_svc.opening_balance(db, school_id=school_id, enrollment=enrollment,
                                     amount_pesewas=pesewas, actor_id=actor_id,
                                     occurred_on=date.today(),
                                     note="Imported opening balance", request=request)
    return student.admission_code


def _import_teacher_row(db: Session, *, school_id: uuid.UUID, norm: dict,
                        year_id: uuid.UUID, actor_id: uuid.UUID,
                        request: Request | None) -> str:
    phone = str(norm.get("phone", "")).strip()
    teacher = Teacher(id=uuid7(), school_id=school_id,
                      staff_code=str(norm.get("staff_code", "")).strip() or None,
                      surname=str(norm["surname"]).strip(),
                      other_names=str(norm["other_names"]).strip(),
                      gender=(GENDER_ALIASES.get(str(norm.get("gender", "")).strip().lower())),
                      phone=normalize_ghana_phone(phone) if phone else None,
                      email=str(norm.get("email", "")).strip() or None,
                      job_title=str(norm.get("job_title", "")).strip() or None,
                      qualification=str(norm.get("qualification", "")).strip() or None,
                      hired_on=_parse_date(norm.get("hired_on")),
                      created_by=actor_id)
    db.add(teacher)
    db.flush()
    audit(db, actor_id=actor_id, action="teacher.imported", entity_type="teacher",
          entity_id=teacher.id, new={"name": teacher.full_name}, request=request)
    # form teacher assignment for primary_class
    cls = str(norm.get("primary_class", "")).strip()
    if cls:
        streams = _streams_by_name(db, school_id, year_id)
        stream = streams.get(cls.lower())
        if stream is not None:
            db.add(TeacherAssignment(id=uuid7(), school_id=school_id, teacher_id=teacher.id,
                                     academic_year_id=year_id, class_stream_id=stream.id,
                                     role="FORM_TEACHER"))
            # subject assignments
            for code in str(norm.get("assigned_subjects", "")).split(";"):
                code = code.strip().upper()
                if not code:
                    continue
                subj = db.scalar(select(Subject).where(Subject.school_id == school_id,
                                                       Subject.code == code))
                if subj is not None:
                    db.add(TeacherAssignment(id=uuid7(), school_id=school_id,
                                             teacher_id=teacher.id, academic_year_id=year_id,
                                             class_stream_id=stream.id, subject_id=subj.id,
                                             role="SUBJECT_TEACHER"))
    return teacher.full_name


def _import_parent_row(db: Session, *, school_id: uuid.UUID, norm: dict,
                       actor_id: uuid.UUID, match_resolution: dict | None,
                       request: Request | None) -> str:
    phone = str(norm.get("phone", "")).strip()
    guardian = None
    if match_resolution and match_resolution.get("action") == "LINK":
        guardian = db.get(ParentGuardian, uuid.UUID(match_resolution["guardian_id"]))
    elif str(norm.get("name", "")).strip():
        guardian = guardians_svc.create_guardian(
            db, school_id=school_id, name=str(norm["name"]).strip(), phone=phone,
            phone2=str(norm.get("phone2", "")).strip() or None,
            email=str(norm.get("email", "")).strip() or None,
            occupation=str(norm.get("occupation", "")).strip() or None,
            residential_address=str(norm.get("residential_address", "")).strip() or None,
            ghana_digital_address=str(norm.get("ghana_digital_address", "")).strip() or None,
            actor_id=actor_id, request=request)
    code = str(norm.get("student_admission_code", "")).strip()
    if guardian is not None and code:
        student = db.scalar(select(Student).where(Student.school_id == school_id,
                                                  Student.admission_code == code))
        if student is not None:
            rel = REL_ALIASES.get(str(norm.get("relationship_to", "")).strip().lower(), "GUARDIAN")
            guardians_svc.link_student(db, school_id=school_id, parent_id=guardian.id,
                                       student_id=student.id, relationship_type=rel,
                                       actor_id=actor_id,
                                       source="IMPORT_SUGGESTION" if match_resolution else "IMPORT",
                                       request=request)
    return guardian.name if guardian else "-"


def commit_job(db: Session, job: ImportJob, *, actor_id: uuid.UUID,
               request: Request | None = None) -> dict:
    """Atomic commit (BR-I01): one transaction; any failure rolls everything back."""
    if job.stage not in ("PREVIEWED",):
        raise ConflictError(f"Job is {job.stage}; only previewed jobs can be confirmed.",
                            code="STAGE_INVALID")
    year_id = uuid.UUID(job.options["academic_year_id"])
    rows = db.scalars(select(ImportErrorRow).where(
        ImportErrorRow.job_id == job.id).order_by(ImportErrorRow.row_no)).all()
    blocked = [r.row_no for r in rows if r.severity == "ERROR"]
    if blocked:
        raise ConflictError(f"{len(blocked)} row(s) have errors — fix and re-upload "
                            f"(rows: {blocked[:10]}{'…' if len(blocked) > 10 else ''}).",
                            code="ERRORS_BLOCK_COMMIT")
    unresolved = [r.row_no for r in rows
                  if r.severity == "MATCH_CANDIDATE"
                  and not (r.match_suggestion or {}).get("resolved")]
    if unresolved:
        raise ConflictError(
            f"{len(unresolved)} parent match(es) need a decision before import "
            f"(rows: {unresolved[:10]}{'…' if len(unresolved) > 10 else ''}). "
            "Phone matches are never auto-linked.", code="MATCHES_UNRESOLVED")
    job.stage = "IMPORTING"
    job.started_at = utcnow()
    db.flush()
    created = skipped_dup = warnings = 0
    details: list[str] = []
    # SAVEPOINT: row-level failures undo all entities but keep the job record,
    # so ROLLED_BACK state is visible afterwards (BR-I01).
    savepoint = db.begin_nested()
    try:
        for r in rows:
            if r.severity == "DUPLICATE":
                skipped_dup += 1
                continue
            resolution = r.match_suggestion if r.severity == "MATCH_CANDIDATE" else None
            if job.kind == "STUDENTS":
                details.append(_import_student_row(
                    db, school_id=job.school_id, norm=r.raw_row, year_id=year_id,
                    actor_id=actor_id, match_resolution=resolution, request=request))
            elif job.kind == "TEACHERS":
                details.append(_import_teacher_row(
                    db, school_id=job.school_id, norm=r.raw_row, year_id=year_id,
                    actor_id=actor_id, request=request))
            else:
                details.append(_import_parent_row(
                    db, school_id=job.school_id, norm=r.raw_row, actor_id=actor_id,
                    match_resolution=resolution, request=request))
            created += 1
            if r.severity == "WARNING":
                warnings += 1
        savepoint.commit() if hasattr(savepoint, "commit") else None
    except Exception as exc:
        savepoint.rollback()
        job.stage = "ROLLED_BACK"
        job.finished_at = utcnow()
        job.summary = {"error": str(exc)[:250]}
        audit(db, actor_id=actor_id, action="import.rolled_back", entity_type="import_job",
              entity_id=job.id, new={"error": str(exc)[:250]}, request=request)
        db.flush()
        return {"rolled_back": True, "error": str(exc)[:250]}
    job.stage = "COMPLETED"
    job.finished_at = utcnow()
    job.summary = {"created": created, "skipped_duplicates": skipped_dup,
                   "warnings": warnings}
    audit(db, actor_id=actor_id, action="import.committed", entity_type="import_job",
          entity_id=job.id, new=job.summary, request=request)
    db.flush()
    return {"rolled_back": False, "created": created,
            "skipped_duplicates": skipped_dup, "warnings": warnings,
            "sample": details[:5]}
